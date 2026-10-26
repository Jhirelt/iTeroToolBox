"""
report.py  -- reporte de escalaciones en Excel, un solo archivo.

Uso:

    from report import make_report

    make_report(
        tickets=[...],
        handlers=[...],
        period="July 2026  -  9 closed escalations",
        title="CSC Escalations  |  July 2026  |  Jhirelt Mora",
        out="reporte.xlsx",
        open_tickets=[...],   # opcional -- ver mas abajo
    )

Formato de los datos
--------------------
tickets      : lista de dicts, cada uno con estas llaves
               ticket, model, ticket_type, issue_type, sub_issue, outcome, csat
               resolution_hours (opcional -- horas habiles, lunes a viernes
               8am-5pm, entre Date/Time Opened y Closed Date; se omite o se
               deja en 0 en el ticket que no lo tenga y ese ticket
               simplemente no cuenta para el promedio, en vez de arrastrarlo
               hacia 0)
handlers     : lista de dicts, cada uno con estas llaves
               ticket, handler, validity, origin
               (validity es "Valid" o "Avoidable")
open_tickets : opcional -- tickets todavia sin cerrar, no llevan outcome/csat
               (no estan resueltos todavia). Lista de dicts, cada uno con
               estas llaves:
               ticket, model, ticket_type, issue_type, sub_issue, comment,
               handler, validity, origin
               (comment es texto libre: el estado actual del issue.
               validity/handler/origin significan lo mismo que en handlers)
               Si se omite o viene vacia, el reporte no incluye la seccion
               de Open Tickets en absoluto.
details      : opcional -- datos de Salesforce que solo van al Excel (no se
               muestran en la app). Lista de dicts, cada uno con estas llaves:
               ticket, status, priority, internal_requestor, date_opened,
               date_closed, days_open
               (date_opened/date_closed pueden ser datetime o texto;
               days_open son dias calendario entre ambas, vacio si sigue
               abierto). Si se omite o viene vacia, no sale la seccion.

Los KPIs y el breakdown se calculan solos a partir de tickets/handlers --
open_tickets nunca entra en esos calculos, son solo para tickets YA
cerrados en el periodo. No hay que escribir ningun numero a mano.
"""

from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------- estilo
# Cambia estos valores y cambian TODOS los reportes.
FONT = "Arial"
NAVY = "1F3864"      # bandas de seccion y titulo
HEAD = "D9E2F3"      # encabezados de tabla
ZEBRA = "F7F9FC"     # filas alternas
ORANGE = "FCE4D6"    # resaltado: reemplazo de parte / avoidable
GREEN = "E2EFDA"     # resaltado: CSAT positivo
MUTED = "808080"

_thin = Side(style="thin", color="C6CEDD")
_B = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

NCOLS = 8
WIDTHS = {"A": 12, "B": 17, "C": 18, "D": 26, "E": 20, "F": 22, "G": 9, "H": 61}


def make_report(tickets, handlers, period, title, out="report.xlsx", open_tickets=None,
                details=None):
    """Genera el workbook y devuelve la ruta."""
    open_tickets = open_tickets or []
    details = details or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    state = {"row": 1}

    # ---------- helpers internos ----------
    def fill(r, c, color):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=color)

    def band(text):
        r = state["row"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
        for c in range(1, NCOLS + 1):
            fill(r, c, NAVY)
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 17
        state["row"] += 1

    def header(labels, span_last=False):
        r = state["row"]
        for i, label in enumerate(labels):
            c = ws.cell(row=r, column=1 + i, value=label)
            c.font = Font(name=FONT, size=9, bold=True, color=NAVY)
            c.fill = PatternFill("solid", fgColor=HEAD)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = _B
        if span_last and len(labels) < NCOLS:
            last = len(labels)
            ws.merge_cells(start_row=r, start_column=last, end_row=r, end_column=NCOLS)
            ws.cell(row=r, column=last, value=labels[-1])
            for c in range(last, NCOLS + 1):
                fill(r, c, HEAD)
                ws.cell(row=r, column=c).border = _B
        ws.row_dimensions[r].height = 24
        state["row"] += 1

    def table(rows, ncols, center=(), height=30, span_last=False,
              highlight=None, fmts=None):
        first = state["row"]
        for i, values in enumerate(rows):
            r = first + i
            for j, v in enumerate(values, start=1):
                ws.cell(row=r, column=j, value=v)
            if span_last and ncols < NCOLS:
                ws.merge_cells(start_row=r, start_column=ncols, end_row=r, end_column=NCOLS)
            ws.row_dimensions[r].height = height
            upto = NCOLS if span_last else ncols
            for c in range(1, upto + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name=FONT, size=9)
                cell.border = _B
                cell.alignment = Alignment(
                    horizontal="center" if c in center else "general",
                    vertical="center", wrap_text=True)
                if i % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor=ZEBRA)
            if fmts:
                for c, f in fmts.items():
                    ws.cell(row=r, column=c).number_format = f
            if highlight:
                for c, color in (highlight(values) or {}).items():
                    fill(r, c, color)
        state["row"] = first + len(rows)

    def caption(text):
        r = state["row"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=8, italic=True, color=MUTED)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 12
        state["row"] += 1

    # ---------- titulo ----------
    ws.merge_cells("A1:H1")
    for c in range(1, NCOLS + 1):
        fill(1, c, NAVY)
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(name=FONT, size=13, bold=True, color="FFFFFF")
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24

    ws.cell(row=2, column=1, value="Closed period").font = Font(name=FONT, size=9, bold=True)
    ws.merge_cells("B2:H2")
    m = ws.cell(row=2, column=2, value=period)
    m.font = Font(name=FONT, size=9, italic=True, color=MUTED)
    m.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[3].height = 5
    state["row"] = 4

    # ---------- 1. key figures (calculados) ----------
    # Only tickets that actually carry a resolution_hours (both Date/Time
    # Opened and Closed Date were read and parsed successfully) count
    # toward the average -- a ticket missing one silently drops out of
    # this figure instead of dragging the average toward 0.
    resolved_hours = [t.get("resolution_hours") or 0 for t in tickets if t.get("resolution_hours")]
    avg_resolution = round(sum(resolved_hours) / len(resolved_hours), 1) if resolved_hours else 0
    kpis = [
        ("Closed", len(tickets)),
        ("Scanner Hardware", sum(1 for t in tickets if t["ticket_type"] == "Scanner Hardware")),
        ("Part replacements", sum(1 for t in tickets if str(t["outcome"]).startswith("Part Replacement"))),
        ("Lumina units", sum(1 for t in tickets if str(t["model"]).startswith("Lumina"))),
        # Counted straight off handlers (same rows the Case Handlers table
        # below lists, one per ticket), not deduped by ticket number — a
        # ticket escalated/closed/re-escalated within the same period can
        # legitimately appear twice, and collapsing to a dict kept only the
        # last one's validity, undercounting both figures.
        ("Valid", sum(1 for h in handlers if h.get("validity") == "Valid")),
        ("Avoidable", sum(1 for h in handlers if h.get("validity") == "Avoidable")),
        ("Positive CSAT", sum(1 for t in tickets if str(t["csat"]).upper() == "YES")),
        ("Avg Resolution (biz hrs)", avg_resolution),
    ]
    band("1.  KEY FIGURES")
    header([k for k, _ in kpis])
    r = state["row"]
    for i, (_, v) in enumerate(kpis):
        c = ws.cell(row=r, column=1 + i, value=v)
        c.font = Font(name=FONT, size=13, bold=True, color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _B
    for c in range(len(kpis) + 1, NCOLS + 1):
        ws.cell(row=r, column=c).border = _B
    ws.row_dimensions[r].height = 24
    state["row"] += 1

    # ---------- 2. closed tickets ----------
    band("2.  CLOSED TICKETS")
    header(["Ticket", "Model", "Ticket Type", "Issue Type", "Sub-Issue Type",
            "Outcome", "CSAT"], span_last=True)
    rows = [(t["ticket"], t["model"], t["ticket_type"], t["issue_type"],
             t.get("sub_issue") or "\u2014", t["outcome"], t["csat"]) for t in tickets]

    def hl_tickets(v):
        out = {}
        if str(v[5]).startswith("Part Replacement"):
            out[6] = ORANGE
        if str(v[6]).upper() == "YES":
            out[7] = GREEN
        return out

    table(rows, 7, center=(1, 5, 7), span_last=True, highlight=hl_tickets)

    # ---------- 3. breakdown (calculado) ----------
    band("3.  BREAKDOWN BY CATEGORY")
    header(["Dimension", "Value", "Tickets", "Share"])
    brk = []
    for label, key in [("Ticket Type", "ticket_type"),
                       ("Issue Type", "issue_type"),
                       ("Sub-Issue Type", "sub_issue")]:
        vals = [t.get(key) for t in tickets if t.get(key)]
        base = len(vals) or 1
        for value, n in Counter(vals).most_common():
            brk.append((label, value, n, n / base))
    table(brk, 4, center=(3, 4), height=16, fmts={4: "0%"})
    missing = sum(1 for t in tickets if not t.get("sub_issue"))
    if missing:
        caption(f"Sub-issue shares are calculated on {len(tickets) - missing} tickets: "
                f"{missing} carry no sub-issue type in the source.")

    # ---------- 4. handlers y origin ----------
    band("4.  CASE HANDLERS AND ORIGIN OF THE ESCALATION")
    header(["Ticket", "Previous Case Handler", "Validity", "Origin of the Escalation"],
           span_last=True)
    hrows = [(h["ticket"], h["handler"], h.get("validity", ""), h.get("origin", ""))
             for h in handlers]
    table(hrows, 4, center=(1, 3), height=44, span_last=True,
          highlight=lambda v: {3: ORANGE} if v[2] == "Avoidable" else {})

    # Numeracion de las secciones opcionales: depende de cuales salen.
    sec = 5

    # ---------- 5. ticket details (opcional) ----------
    # Priority / Internal Requestor / fechas de Salesforce -- solo Excel,
    # la app nunca los muestra. Incluye abiertos y cerrados.
    if details:
        band(f"{sec}.  TICKET DETAILS")
        sec += 1
        header(["Ticket", "Status", "Priority", "Internal Requestor",
                "Date/Time Opened", "Date/Time Closed", "Days Open"], span_last=True)
        drows = [(d.get("ticket", ""), d.get("status", ""), d.get("priority") or "—",
                  d.get("internal_requestor") or "—",
                  d.get("date_opened") or "—", d.get("date_closed") or "—",
                  d.get("days_open") if d.get("days_open") not in (None, "") else "—")
                 for d in details]
        dt_fmt = "mm/dd/yyyy h:mm AM/PM"
        table(drows, 7, center=(1, 2, 3, 5, 6, 7), height=16, span_last=True,
              fmts={5: dt_fmt, 6: dt_fmt})

    # ---------- open tickets (opcional -- solo si hay alguno) -------
    # No entran en Key Figures/Breakdown de arriba -- esas secciones son
    # sobre lo YA cerrado en el periodo, un ticket todavia abierto no
    # cuenta ahi. Mismo formato de dos tablas que Closed Tickets +
    # Handlers arriba, solo que "Outcome/CSAT" se reemplaza por un
    # Comment de texto libre (no hay resultado final todavia).
    if open_tickets:
        band(f"{sec}.  OPEN TICKETS")
        header(["Ticket", "Model", "Ticket Type", "Issue Type", "Sub-Issue Type",
                "Comment — Current Status"], span_last=True)
        orows = [(o["ticket"], o["model"], o["ticket_type"], o["issue_type"],
                  o.get("sub_issue") or "—", o.get("comment") or "—")
                 for o in open_tickets]
        table(orows, 6, center=(1, 5), height=44, span_last=True)

        band(f"{sec + 1}.  OPEN TICKETS — CASE HANDLERS AND ORIGIN OF THE ESCALATION")
        header(["Ticket", "Previous Case Handler", "Validity", "Origin of the Escalation"],
               span_last=True)
        ohrows = [(o["ticket"], o.get("handler", ""), o.get("validity", ""), o.get("origin", ""))
                  for o in open_tickets]
        table(ohrows, 4, center=(1, 3), height=44, span_last=True,
              highlight=lambda v: {3: ORANGE} if v[2] == "Avoidable" else {})

    ws.freeze_panes = "A4"
    wb.save(out)
    return out


# ------------------------------------------------------------------ demo
if __name__ == "__main__":
    tickets = [
        dict(ticket="67306304", model="Lumina Pro - Cart", ticket_type="Network",
             issue_type="Ports / Firewall / Restriction Issue", sub_issue="",
             outcome="Network Troubleshooting", csat="NO", resolution_hours=4.8),
        dict(ticket="67398715", model="Lumina - Cart", ticket_type="Scanner Hardware",
             issue_type="Lumina Cart Configuration", sub_issue="Wand disconnections",
             outcome="Part Replacement \u2013 Wand", csat="NO", resolution_hours=26.25),
        dict(ticket="67663575", model="Lumina Pro - Cart", ticket_type="Scanner Hardware",
             issue_type="Lumina Cart Configuration", sub_issue="Customer satisfaction",
             outcome="Part Replacement \u2013 PPC", csat="YES", resolution_hours=9.0),
    ]
    handlers = [
        dict(ticket="67306304", handler="Jose Fabio Granados / T3", validity="Valid",
             origin="Issue identified on the first call. The office had an outdated port list."),
        dict(ticket="67398715", handler="Adrian Quesada", validity="Valid",
             origin="Escalated after two failed attempts to fix the issue."),
        dict(ticket="67663575", handler="Aaron Raul Ramirez Loria", validity="Avoidable",
             origin="Lack of follow-up originated the case getting escalated."),
    ]
    open_tickets = [
        dict(ticket="67701122", model="Lumina - Mobile", ticket_type="Software",
             issue_type="Design Suite", sub_issue="",
             comment="Waiting on customer to reproduce with screen recording.",
             handler="Priscilla Rodriguez", validity="Valid",
             origin="Escalated after repeated crashes with no clear repro steps."),
    ]
    print(make_report(tickets, handlers,
                      period="Demo  -  3 closed escalations",
                      title="CSC Escalations  |  Demo",
                      out="demo.xlsx",
                      open_tickets=open_tickets))
